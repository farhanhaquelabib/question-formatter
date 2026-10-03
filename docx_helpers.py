import os
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    """Set inner margins for a table cell in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set borders for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{top or "none"}" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:bottom w:val="{bottom or "none"}" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:left w:val="{left or "none"}" w:sz="4" w:space="0" w:color="000000"/>'
        f'<w:right w:val="{right or "none"}" w:sz="4" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def set_cell_shading(cell, color_hex):
    """Set background color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color_hex}"/>')
    tcPr.append(shd)

def add_english_run(p, text, bold=False, italic=False, size_pt=11, color_rgb=None):
    run = p.add_run(text)
    run.bold = bold
    run.italic = False # Strict rule 1: No italic words
    if color_rgb:
        run.font.color.rgb = color_rgb
    run.font.name = 'Cambria Math'
    run.font.size = Pt(size_pt)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>')
    rPr.append(rFonts)
    # Explicitly enforce no italics in Word XML
    i_xml = parse_xml(f'<w:i {nsdecls("w")} w:val="0"/>')
    iCs_xml = parse_xml(f'<w:iCs {nsdecls("w")} w:val="0"/>')
    rPr.append(i_xml)
    rPr.append(iCs_xml)
    sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(size_pt * 2)}"/>')
    rPr.append(sz)
    return run

def add_bangla_run(p, text, bold=False, italic=False, size_pt=9, color_rgb=None):
    run = p.add_run(text)
    run.bold = bold
    run.italic = False # Strict rule 1: No italic words
    if color_rgb:
        run.font.color.rgb = color_rgb
    run.font.name = 'Nirmala UI'
    run.font.size = Pt(size_pt)
    rPr = run._r.get_or_add_rPr()
    rFonts = parse_xml(f'<w:rFonts {nsdecls("w")} w:ascii="Nirmala UI" w:hAnsi="Nirmala UI" w:cs="Nirmala UI"/>')
    rPr.append(rFonts)
    # Explicitly enforce no italics in Word XML
    i_xml = parse_xml(f'<w:i {nsdecls("w")} w:val="0"/>')
    iCs_xml = parse_xml(f'<w:iCs {nsdecls("w")} w:val="0"/>')
    rPr.append(i_xml)
    rPr.append(iCs_xml)
    sz = parse_xml(f'<w:sz {nsdecls("w")} w:val="{int(size_pt * 2)}"/>')
    szCs = parse_xml(f'<w:szCs {nsdecls("w")} w:val="{int(size_pt * 2)}"/>')
    rPr.append(sz)
    rPr.append(szCs)
    return run

def _build_omml_element(content, bold=False):
    """Build OMML XML fragment for text, radical, fraction, superscript, or nested expression."""
    bold_tag = "<w:b/>" if bold else ""
    if isinstance(content, tuple):
        op = content[0]
        if op == "sqrt":
            radicand = content[1]
            rad_inner = _build_omml_element(radicand, bold=bold)
            return (
                f'<m:rad>'
                f'<m:radPr><m:degHide m:val="1"/></m:radPr>'
                f'<m:deg/>'
                f'<m:e>{rad_inner}</m:e>'
                f'</m:rad>'
            )
        elif op == "frac":
            num_x = _build_omml_element(content[1], bold=bold)
            den_x = _build_omml_element(content[2], bold=bold)
            return (
                f'<m:f>'
                f'<m:fPr><m:type m:val="bar"/></m:fPr>'
                f'<m:num>{num_x}</m:num>'
                f'<m:den>{den_x}</m:den>'
                f'</m:f>'
            )
        elif op == "sup":
            base_x = _build_omml_element(content[1], bold=bold)
            sup_x = _build_omml_element(content[2], bold=bold)
            return (
                f'<m:sSup>'
                f'<m:e>{base_x}</m:e>'
                f'<m:sup>{sup_x}</m:sup>'
                f'</m:sSup>'
            )
    elif isinstance(content, str):
        if content.startswith("√"):
            radicand = content[1:]
            return (
                f'<m:rad>'
                f'<m:radPr><m:degHide m:val="1"/></m:radPr>'
                f'<m:deg/>'
                f'<m:e><m:r><m:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><m:nor/>{bold_tag}</m:rPr><m:t>{radicand}</m:t></m:r></m:e>'
                f'</m:rad>'
            )
        elif content.startswith("sqrt(") and content.endswith(")"):
            radicand = content[5:-1]
            return (
                f'<m:rad>'
                f'<m:radPr><m:degHide m:val="1"/></m:radPr>'
                f'<m:deg/>'
                f'<m:e><m:r><m:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><m:nor/>{bold_tag}</m:rPr><m:t>{radicand}</m:t></m:r></m:e>'
                f'</m:rad>'
            )
        else:
            return f'<m:r><m:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><m:nor/>{bold_tag}</m:rPr><m:t>{content}</m:t></m:r>'
    return f'<m:r><m:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><m:nor/>{bold_tag}</m:rPr><m:t>{str(content)}</m:t></m:r>'

def add_omml_fraction(paragraph, num_content, den_content, bold=False):
    """Add a native Office Math (OMML) proper vertical fraction with non-italic math."""
    num_xml = _build_omml_element(num_content, bold=bold)
    den_xml = _build_omml_element(den_content, bold=bold)
    omml_str = (
        f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<m:f>'
        f'<m:fPr><m:type m:val="bar"/></m:fPr>'
        f'<m:num>{num_xml}</m:num>'
        f'<m:den>{den_xml}</m:den>'
        f'</m:f>'
        f'</m:oMath>'
    )
    paragraph._p.append(parse_xml(omml_str))

def add_omml_sqrt(paragraph, expr_text, bold=False):
    """Add a native Office Math (OMML) square root with vincular bar spanning over expr_text."""
    bold_tag = "<w:b/>" if bold else ""
    omml_str = (
        f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<m:rad>'
        f'<m:radPr><m:degHide m:val="1"/></m:radPr>'
        f'<m:deg/>'
        f'<m:e><m:r><m:rPr><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/><m:nor/>{bold_tag}</m:rPr><m:t>{expr_text}</m:t></m:r></m:e>'
        f'</m:rad>'
        f'</m:oMath>'
    )
    paragraph._p.append(parse_xml(omml_str))

def add_omml_sup(paragraph, base, exp, bold=False):
    """Add a native Office Math (OMML) superscript expression base^exp."""
    base_xml = _build_omml_element(base, bold=bold)
    exp_xml = _build_omml_element(exp, bold=bold)
    omml_str = (
        f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        f'<m:sSup>'
        f'<m:e>{base_xml}</m:e>'
        f'<m:sup>{exp_xml}</m:sup>'
        f'</m:sSup>'
        f'</m:oMath>'
    )
    paragraph._p.append(parse_xml(omml_str))

def add_omml_paren_frac(paragraph, num_content, den_content, bold=False, sup=None):
    """
    Add a fraction enclosed inside auto-scaling delimiter brackets ( ) via <m:d>,
    optionally raised to an exponent / superscript via <m:sSup>.
    """
    num_xml = _build_omml_element(num_content, bold=bold)
    den_xml = _build_omml_element(den_content, bold=bold)
    bold_tag = "<w:b/>" if bold else ""
    
    frac_xml = (
        f'<m:f>'
        f'<m:fPr><m:type m:val="bar"/></m:fPr>'
        f'<m:num>{num_xml}</m:num>'
        f'<m:den>{den_xml}</m:den>'
        f'</m:f>'
    )
    delim_xml = (
        f'<m:d>'
        f'<m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
        f'<m:e>{frac_xml}</m:e>'
        f'</m:d>'
    )
    if sup:
        sup_xml = _build_omml_element(sup, bold=bold)
        omml_str = (
            f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<m:sSup>'
            f'<m:e>{delim_xml}</m:e>'
            f'<m:sup>{sup_xml}</m:sup>'
            f'</m:sSup>'
            f'</m:oMath>'
        )
    else:
        omml_str = (
            f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'{delim_xml}'
            f'</m:oMath>'
        )
    paragraph._p.append(parse_xml(omml_str))

def add_omml_paren_expr(paragraph, expr_items, bold=False, sup=None):
    """
    Add an expression enclosed inside auto-scaling delimiter brackets ( ),
    optionally raised to an exponent. E.g. (x + 1/x)^2
    """
    inner_xml = "".join(_build_omml_element(it, bold=bold) for it in expr_items)
    delim_xml = (
        f'<m:d>'
        f'<m:dPr><m:begChr m:val="("/><m:endChr m:val=")"/></m:dPr>'
        f'<m:e>{inner_xml}</m:e>'
        f'</m:d>'
    )
    if sup:
        sup_xml = _build_omml_element(sup, bold=bold)
        omml_str = (
            f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'<m:sSup>'
            f'<m:e>{delim_xml}</m:e>'
            f'<m:sup>{sup_xml}</m:sup>'
            f'</m:sSup>'
            f'</m:oMath>'
        )
    else:
        omml_str = (
            f'<m:oMath xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f'{delim_xml}'
            f'</m:oMath>'
        )
    paragraph._p.append(parse_xml(omml_str))

def add_math_segments(p, segments, is_bangla=False, size_pt=11, bold=False):
    """
    Universal segment handler for English, Bangla, and Math equations.
    Segments can contain:
      - str: plain text
      - ("frac", num, den): proper vertical fraction
      - ("sqrt", expr): radical square root
      - ("sup", base, exp): superscript
      - ("paren_frac", num, den): auto-scaling bracketed fraction (num/den)
      - ("paren_frac_sup", num, den, sup): bracketed fraction with exponent (num/den)^sup
      - ("paren_expr_sup", items, sup): bracketed expression with exponent (expr)^sup
    """
    for item in segments:
        if isinstance(item, str):
            if is_bangla:
                add_bangla_run(p, item, bold=False, size_pt=size_pt)
            else:
                add_english_run(p, item, bold=bold, size_pt=size_pt)
        elif isinstance(item, tuple):
            tag = item[0]
            if tag == "frac":
                add_omml_fraction(p, item[1], item[2], bold=bold)
            elif tag == "sqrt":
                add_omml_sqrt(p, item[1], bold=bold)
            elif tag == "sup":
                add_omml_sup(p, item[1], item[2], bold=bold)
            elif tag == "paren_frac":
                add_omml_paren_frac(p, item[1], item[2], bold=bold)
            elif tag == "paren_frac_sup":
                add_omml_paren_frac(p, item[1], item[2], bold=bold, sup=item[3])
            elif tag == "paren_expr_sup":
                add_omml_paren_expr(p, item[1], bold=bold, sup=item[2])

def add_english_math_segments(p, segments, size_pt=11, bold=False):
    add_math_segments(p, segments, is_bangla=False, size_pt=size_pt, bold=bold)

def add_bangla_math_segments(p, segments, size_pt=9):
    add_math_segments(p, segments, is_bangla=True, size_pt=size_pt, bold=False)

# 4 columns on Letter page (8.5" x 11.0") with 0.75" margins (7.0" printable width = 10080 dxa):
# Col 1: 0" (0 dxa) | Col 2: 1.75" (2520 dxa) | Col 3: 3.50" (5040 dxa) | Col 4: 5.25" (7560 dxa)
TAB_STOPS_075 = [2520, 5040, 7560]

def set_exam_page_margins(section, top=0.75, bottom=0.75, left=0.75, right=0.75):
    """Set standard exam margins in inches (default 0.75" on all sides)."""
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)

def set_paragraph_tab_stops(p, positions_dxa=TAB_STOPS_075):
    """
    Add ruler tab stops to paragraph for multi-column alignment without tables.
    positions_dxa: list of positions in dxa (1 inch = 1440 dxa). Default is [2520, 5040, 7560].
    """
    pPr = p._p.get_or_add_pPr()
    tabs = parse_xml(f'<w:tabs {nsdecls("w")}/>')
    for pos in positions_dxa:
        tab_elem = parse_xml(f'<w:tab {nsdecls("w")} w:val="left" w:pos="{pos}"/>')
        tabs.append(tab_elem)
    pPr.append(tabs)

def add_tab_run(p):
    """
    Add a tab run to advance to the next ruler tab stop.
    """
    run = p.add_run()
    run.element.append(parse_xml(f'<w:tab {nsdecls("w")}/>'))
    return run

print("Loaded docx_helpers with Phoenix 0.75 margin support, zero italics, auto-scaling brackets, and ruler tab stops.")
